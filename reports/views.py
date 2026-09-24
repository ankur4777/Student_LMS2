from decimal import Decimal
import csv
from io import BytesIO


from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from academics.models import AcademicSession, ClassRoom, Section, StudentEnrollment, Subject
from accounts.models import StudentProfile, TeacherProfile
from assignments.models import Assignment, AssignmentSubmission
from attendance.models import StudentAttendance
from fees.models import FeePayment, StudentFee
from liveclasses.models import LiveClass
from studentresults.models import Exam, StudentResult

ZERO = Decimal("0.00")


def college_admin_organization(user):
    if user.role != "college_admin" or not user.is_active:
        return None
    return user.organization


def decimal_or_zero(value):
    return value or ZERO


def attendance_summary(queryset):
    counts = queryset.aggregate(
        present=Count("id", filter=Q(status=StudentAttendance.Status.PRESENT)),
        absent=Count("id", filter=Q(status=StudentAttendance.Status.ABSENT)),
        late=Count("id", filter=Q(status=StudentAttendance.Status.LATE)),
        excused=Count("id", filter=Q(status=StudentAttendance.Status.EXCUSED)),
    )
    counted = counts["present"] + counts["absent"] + counts["late"]
    counts["overall_percentage"] = (
        round(((counts["present"] + counts["late"]) / counted) * 100, 2)
        if counted else 0
    )
    return counts


def parse_filters(request, organization):
    params = request.query_params
    result = {"session": None, "classroom": None, "section": None, "subject": None}
    model_map = [
        ("session", "academic_session", AcademicSession),
        ("classroom", "classroom", ClassRoom),
        ("section", "section", Section),
        ("subject", "subject", Subject),
    ]
    for key, param, model in model_map:
        raw = params.get(param)
        if raw:
            try:
                result[key] = model.objects.get(pk=raw, organization=organization)
            except (model.DoesNotExist, ValueError):
                return None, f"Invalid {param}."
    if result["classroom"] and result["session"] and result["classroom"].academic_session_id != result["session"].id:
        return None, "Class does not belong to the selected academic session."
    if result["section"] and result["classroom"] and result["section"].classroom_id != result["classroom"].id:
        return None, "Section does not belong to the selected class."
    if result["subject"] and result["classroom"] and result["subject"].classroom_id != result["classroom"].id:
        return None, "Subject does not belong to the selected class."
    start = parse_date(params.get("date_from", "")) if params.get("date_from") else None
    end = parse_date(params.get("date_to", "")) if params.get("date_to") else None
    if params.get("date_from") and not start:
        return None, "Invalid date_from."
    if params.get("date_to") and not end:
        return None, "Invalid date_to."
    if start and end and start > end:
        return None, "date_from cannot be after date_to."
    result.update({"date_from": start, "date_to": end})
    return result, None


def scoped_data(organization, filters):
    attendance = StudentAttendance.objects.filter(attendance_session__organization=organization)
    assignments = Assignment.objects.filter(organization=organization)
    submissions = AssignmentSubmission.objects.filter(assignment__organization=organization)
    live = LiveClass.objects.filter(organization=organization)
    fees = StudentFee.objects.filter(organization=organization)
    payments = FeePayment.objects.filter(organization=organization)
    exams = Exam.objects.filter(organization=organization, is_published=True)
    results = StudentResult.objects.filter(exam__organization=organization, exam__is_published=True)

    session, classroom, section, subject = (
        filters["session"], filters["classroom"], filters["section"], filters["subject"]
    )
    if session:
        attendance = attendance.filter(attendance_session__section__classroom__academic_session=session)
        assignments = assignments.filter(teacher_assignment__section__classroom__academic_session=session)
        submissions = submissions.filter(assignment__teacher_assignment__section__classroom__academic_session=session)
        live = live.filter(teacher_assignment__section__classroom__academic_session=session)
        fees = fees.filter(academic_session=session)
        payments = payments.filter(student_fee__academic_session=session)
        exams = exams.filter(section__classroom__academic_session=session)
        results = results.filter(exam__section__classroom__academic_session=session)
    if classroom:
        attendance = attendance.filter(attendance_session__section__classroom=classroom)
        assignments = assignments.filter(teacher_assignment__section__classroom=classroom)
        submissions = submissions.filter(assignment__teacher_assignment__section__classroom=classroom)
        live = live.filter(teacher_assignment__section__classroom=classroom)
        fees = fees.filter(fee_structure__class_room=classroom)
        payments = payments.filter(student_fee__fee_structure__class_room=classroom)
        exams = exams.filter(section__classroom=classroom)
        results = results.filter(exam__section__classroom=classroom)
    if section:
        attendance = attendance.filter(attendance_session__section=section)
        assignments = assignments.filter(teacher_assignment__section=section)
        submissions = submissions.filter(assignment__teacher_assignment__section=section)
        live = live.filter(teacher_assignment__section=section)
        exams = exams.filter(section=section)
        results = results.filter(exam__section=section)
        student_ids = StudentEnrollment.objects.filter(section=section, is_active=True).values("student_id")
        fees = fees.filter(student_id__in=student_ids)
        payments = payments.filter(student_fee__student_id__in=student_ids)
    if subject:
        attendance = attendance.filter(attendance_session__subject=subject)
        assignments = assignments.filter(teacher_assignment__subject=subject)
        submissions = submissions.filter(assignment__teacher_assignment__subject=subject)
        live = live.filter(teacher_assignment__subject=subject)
        results = results.filter(subject=subject)
    if filters["date_from"]:
        attendance = attendance.filter(attendance_session__date__gte=filters["date_from"])
        assignments = assignments.filter(due_date__gte=filters["date_from"])
        live = live.filter(class_date__gte=filters["date_from"])
        payments = payments.filter(payment_date__gte=filters["date_from"])
        exams = exams.filter(exam_date__gte=filters["date_from"])
        results = results.filter(exam__exam_date__gte=filters["date_from"])
    if filters["date_to"]:
        attendance = attendance.filter(attendance_session__date__lte=filters["date_to"])
        assignments = assignments.filter(due_date__lte=filters["date_to"])
        live = live.filter(class_date__lte=filters["date_to"])
        payments = payments.filter(payment_date__lte=filters["date_to"])
        exams = exams.filter(exam_date__lte=filters["date_to"])
        results = results.filter(exam__exam_date__lte=filters["date_to"])
    return attendance, assignments, submissions, live, fees, payments, exams, results


class CollegeAdminFilterOptionsAPIView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        org = college_admin_organization(request.user)
        if not org:
            return Response({"detail": "College admin access required."}, status=403)
        return Response({
            "academic_sessions": list(AcademicSession.objects.filter(organization=org).values("id", "name", "is_active")),
            "classes": list(ClassRoom.objects.filter(organization=org).values("id", "name", "academic_session_id")),
            "sections": list(Section.objects.filter(organization=org).values("id", "name", "classroom_id")),
            "subjects": list(Subject.objects.filter(organization=org).values("id", "name", "code", "classroom_id")),
        })


class CollegeAdminOverviewAPIView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        org = college_admin_organization(request.user)
        if not org:
            return Response({"detail": "College admin access required."}, status=403)
        filters, error = parse_filters(request, org)
        if error:
            return Response({"detail": error}, status=400)
        attendance, assignments, submissions, live, fees, payments, exams, results = scoped_data(org, filters)
        att = attendance_summary(attendance)
        expected_submissions = 0
        for assignment in assignments.select_related("teacher_assignment__section"):
            expected_submissions += StudentEnrollment.objects.filter(
                section=assignment.teacher_assignment.section, is_active=True,
                student__user__organization=org,
            ).count()
        expected = decimal_or_zero(fees.aggregate(total=Sum("payable_amount"))["total"])
        collected = decimal_or_zero(payments.aggregate(total=Sum("amount"))["total"])
        classroom_qs = ClassRoom.objects.filter(organization=org)
        section_qs = Section.objects.filter(organization=org)
        subject_qs = Subject.objects.filter(organization=org)
        student_qs = StudentProfile.objects.filter(user__organization=org, user__role="student")
        if filters["session"]:
            classroom_qs = classroom_qs.filter(academic_session=filters["session"])
            section_qs = section_qs.filter(classroom__academic_session=filters["session"])
            subject_qs = subject_qs.filter(classroom__academic_session=filters["session"])
            student_qs = student_qs.filter(enrollments__section__classroom__academic_session=filters["session"], enrollments__is_active=True)
        if filters["classroom"]:
            classroom_qs = classroom_qs.filter(pk=filters["classroom"].pk)
            section_qs = section_qs.filter(classroom=filters["classroom"])
            subject_qs = subject_qs.filter(classroom=filters["classroom"])
            student_qs = student_qs.filter(enrollments__section__classroom=filters["classroom"], enrollments__is_active=True)
        if filters["section"]:
            section_qs = section_qs.filter(pk=filters["section"].pk)
            student_qs = student_qs.filter(enrollments__section=filters["section"], enrollments__is_active=True)
        live_counts = live.aggregate(
            scheduled=Count("id", filter=Q(status=LiveClass.Status.SCHEDULED)),
            completed=Count("id", filter=Q(status=LiveClass.Status.COMPLETED)),
            cancelled=Count("id", filter=Q(status=LiveClass.Status.CANCELLED)),
        )
        return Response({
            "academic": {
                "total_students": student_qs.distinct().count(),
                "total_teachers": TeacherProfile.objects.filter(user__organization=org, user__role="teacher").count(),
                "total_classes": classroom_qs.count(), "total_sections": section_qs.count(), "total_subjects": subject_qs.count(),
            },
            "attendance": att,
            "assignments": {"total_assignments": assignments.count(), "total_submissions": submissions.count(),
                "expected_submissions": expected_submissions, "pending_submissions": max(expected_submissions-submissions.count(), 0)},
            "live_classes": {"total_live_classes": live.count(), **live_counts},
            "fees": {"total_expected_amount": expected, "total_collected_amount": collected,
                "total_pending_amount": max(expected-collected, ZERO),
                "paid_count": fees.filter(status=StudentFee.Status.PAID).count(),
                "pending_count": fees.filter(status__in=[StudentFee.Status.PENDING, StudentFee.Status.PARTIALLY_PAID, StudentFee.Status.OVERDUE]).count()},
            "results": {"published_exams": exams.count(), "published_results": results.count()},
        })


class CollegeAdminDetailedAnalyticsAPIView(APIView):
    permission_classes = [IsAuthenticated]
    def get(self, request):
        org = college_admin_organization(request.user)
        if not org:
            return Response({"detail": "College admin access required."}, status=403)
        filters, error = parse_filters(request, org)
        if error:
            return Response({"detail": error}, status=400)
        attendance, assignments, submissions, live, fees, payments, exams, results = scoped_data(org, filters)

        class_qs = ClassRoom.objects.filter(organization=org).select_related("academic_session")
        if filters["session"]: class_qs = class_qs.filter(academic_session=filters["session"])
        if filters["classroom"]: class_qs = class_qs.filter(pk=filters["classroom"].pk)
        class_rows = []
        for classroom in class_qs:
            enrollments = StudentEnrollment.objects.filter(section__classroom=classroom, is_active=True)
            ca = attendance.filter(attendance_session__section__classroom=classroom)
            cf = fees.filter(student_id__in=enrollments.values("student_id"))
            cp = payments.filter(student_fee__student_id__in=enrollments.values("student_id"))
            expected = decimal_or_zero(cf.aggregate(t=Sum("payable_amount"))["t"])
            collected = decimal_or_zero(cp.aggregate(t=Sum("amount"))["t"])
            class_rows.append({"id": classroom.id, "name": classroom.name, "academic_session": classroom.academic_session.name,
                "students": enrollments.count(), "sections": Section.objects.filter(classroom=classroom).count(),
                "attendance_percentage": attendance_summary(ca)["overall_percentage"],
                "assignments": assignments.filter(teacher_assignment__section__classroom=classroom).count(),
                "published_results": results.filter(exam__section__classroom=classroom).count(),
                "expected_fees": expected, "collected_fees": collected, "pending_fees": max(expected-collected, ZERO)})

        low_students = []
        student_ids = attendance.values_list("student_id", flat=True).distinct()
        for enrollment in StudentEnrollment.objects.filter(student_id__in=student_ids, is_active=True).select_related("student__user", "section__classroom"):
            sa = attendance.filter(student=enrollment.student)
            summary = attendance_summary(sa)
            if summary["overall_percentage"] < 75:
                user = enrollment.student.user
                low_students.append({"student": user.get_full_name() or user.username, "roll_number": enrollment.roll_number,
                    "class": enrollment.section.classroom.name, "section": enrollment.section.name,
                    "present": summary["present"], "absent": summary["absent"], "late": summary["late"],
                    "attendance_percentage": summary["overall_percentage"]})

        outstanding = []
        for fee in fees.select_related("student__user", "fee_structure__class_room").prefetch_related("payments"):
            paid = fee.paid_amount
            pending = max(fee.payable_amount-paid, ZERO)
            if pending > ZERO:
                enrollment = StudentEnrollment.objects.filter(student=fee.student, is_active=True).select_related("section__classroom").first()
                outstanding.append({"student": fee.student.user.get_full_name() or fee.student.user.username,
                    "roll_number": enrollment.roll_number if enrollment else "", "class": enrollment.section.classroom.name if enrollment else "",
                    "section": enrollment.section.name if enrollment else "", "expected": fee.payable_amount,
                    "paid": paid, "pending": pending, "status": fee.calculated_status()})

        section_rows = []
        section_qs = Section.objects.filter(organization=org).select_related("classroom")
        if filters["classroom"]: section_qs = section_qs.filter(classroom=filters["classroom"])
        if filters["section"]: section_qs = section_qs.filter(pk=filters["section"].pk)
        for sec in section_qs:
            section_rows.append({"id": sec.id, "class": sec.classroom.name, "section": sec.name,
                "students": StudentEnrollment.objects.filter(section=sec, is_active=True).count(),
                "attendance_percentage": attendance_summary(attendance.filter(attendance_session__section=sec))["overall_percentage"]})

        subject_rows = []
        subject_qs = Subject.objects.filter(organization=org)
        if filters["classroom"]: subject_qs = subject_qs.filter(classroom=filters["classroom"])
        if filters["subject"]: subject_qs = subject_qs.filter(pk=filters["subject"].pk)
        for sub in subject_qs:
            sr = results.filter(subject=sub)
            total_obtained = decimal_or_zero(sr.aggregate(t=Sum("marks_obtained"))["t"])
            total_max = decimal_or_zero(sr.aggregate(t=Sum("maximum_marks"))["t"])
            avg = round(float(total_obtained / total_max * 100), 2) if total_max else 0
            subject_rows.append({"id": sub.id, "subject": sub.name,
                "attendance_percentage": attendance_summary(attendance.filter(attendance_session__subject=sub))["overall_percentage"],
                "assignments": assignments.filter(teacher_assignment__subject=sub).count(),
                "published_results": sr.count(), "average_percentage": avg})

        teacher_rows = []
        teacher_ids = live.values_list("teacher_assignment__teacher_id", flat=True).distinct()
        for teacher in TeacherProfile.objects.filter(pk__in=teacher_ids).select_related("user"):
            tl = live.filter(teacher_assignment__teacher=teacher)
            teacher_rows.append({"teacher": teacher.user.get_full_name() or teacher.user.username,
                "scheduled": tl.filter(status=LiveClass.Status.SCHEDULED).count(),
                "completed": tl.filter(status=LiveClass.Status.COMPLETED).count(),
                "cancelled": tl.filter(status=LiveClass.Status.CANCELLED).count()})

        return Response({"classes": class_rows, "sections": section_rows, "subjects": subject_rows,
            "low_attendance_students": low_students[:50], "outstanding_fees": outstanding[:50],
            "teacher_activity": teacher_rows})


class CollegeAdminCSVExportAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        org = college_admin_organization(request.user)
        if not org:
            return Response({"detail": "College admin access required."}, status=403)

        filters, error = parse_filters(request, org)
        if error:
            return Response({"detail": error}, status=400)

        attendance, assignments, submissions, live, fees, payments, exams, results = scoped_data(
            org, filters
        )
        att = attendance_summary(attendance)
        expected = decimal_or_zero(fees.aggregate(total=Sum("payable_amount"))["total"])
        collected = decimal_or_zero(payments.aggregate(total=Sum("amount"))["total"])

        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="college-admin-report.csv"'
        response.write("\ufeff")
        writer = csv.writer(response)

        writer.writerow(["College Admin Reports & Analytics"])
        writer.writerow(["Organization", org.name])
        writer.writerow([])

        writer.writerow(["Applied Filters"])
        writer.writerow(["Academic Session", filters["session"].name if filters["session"] else "All"])
        writer.writerow(["Class", filters["classroom"].name if filters["classroom"] else "All"])
        writer.writerow(["Section", filters["section"].name if filters["section"] else "All"])
        writer.writerow(["Subject", filters["subject"].name if filters["subject"] else "All"])
        writer.writerow(["Date From", filters["date_from"] or "All"])
        writer.writerow(["Date To", filters["date_to"] or "All"])
        writer.writerow([])

        writer.writerow(["Overview"])
        writer.writerow(["Metric", "Value"])
        writer.writerow(["Overall Attendance %", att["overall_percentage"]])
        writer.writerow(["Present", att["present"]])
        writer.writerow(["Absent", att["absent"]])
        writer.writerow(["Late", att["late"]])
        writer.writerow(["Excused", att["excused"]])
        writer.writerow(["Total Assignments", assignments.count()])
        writer.writerow(["Total Submissions", submissions.count()])
        writer.writerow(["Total Live Classes", live.count()])
        writer.writerow(["Published Exams", exams.count()])
        writer.writerow(["Published Results", results.count()])
        writer.writerow(["Expected Fees", expected])
        writer.writerow(["Collected Fees", collected])
        writer.writerow(["Pending Fees", max(expected - collected, ZERO)])
        writer.writerow([])

        writer.writerow(["Class-wise Performance"])
        writer.writerow([
            "Class", "Session", "Students", "Sections", "Attendance %",
            "Assignments", "Published Results", "Expected Fees",
            "Collected Fees", "Pending Fees",
        ])

        class_qs = ClassRoom.objects.filter(organization=org).select_related("academic_session")
        if filters["session"]:
            class_qs = class_qs.filter(academic_session=filters["session"])
        if filters["classroom"]:
            class_qs = class_qs.filter(pk=filters["classroom"].pk)

        for classroom in class_qs:
            enrollments = StudentEnrollment.objects.filter(
                section__classroom=classroom,
                is_active=True,
            )
            class_attendance = attendance.filter(
                attendance_session__section__classroom=classroom
            )
            class_fees = fees.filter(student_id__in=enrollments.values("student_id"))
            class_payments = payments.filter(
                student_fee__student_id__in=enrollments.values("student_id")
            )
            class_expected = decimal_or_zero(
                class_fees.aggregate(total=Sum("payable_amount"))["total"]
            )
            class_collected = decimal_or_zero(
                class_payments.aggregate(total=Sum("amount"))["total"]
            )
            writer.writerow([
                classroom.name,
                classroom.academic_session.name,
                enrollments.count(),
                Section.objects.filter(classroom=classroom).count(),
                attendance_summary(class_attendance)["overall_percentage"],
                assignments.filter(
                    teacher_assignment__section__classroom=classroom
                ).count(),
                results.filter(exam__section__classroom=classroom).count(),
                class_expected,
                class_collected,
                max(class_expected - class_collected, ZERO),
            ])

        writer.writerow([])
        writer.writerow(["Outstanding Fees"])
        writer.writerow([
            "Student", "Roll Number", "Class", "Section",
            "Expected", "Paid", "Pending", "Status",
        ])
        for fee in fees.select_related("student__user").prefetch_related("payments"):
            paid = fee.paid_amount
            pending = max(fee.payable_amount - paid, ZERO)
            if pending <= ZERO:
                continue
            enrollment = StudentEnrollment.objects.filter(
                student=fee.student,
                is_active=True,
            ).select_related("section__classroom").first()
            writer.writerow([
                fee.student.user.get_full_name() or fee.student.user.username,
                enrollment.roll_number if enrollment else "",
                enrollment.section.classroom.name if enrollment else "",
                enrollment.section.name if enrollment else "",
                fee.payable_amount,
                paid,
                pending,
                fee.calculated_status(),
            ])

        return response


class CollegeAdminExcelExportAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        org = college_admin_organization(request.user)
        if not org:
            return Response({"detail": "College admin access required."}, status=403)

        filters, error = parse_filters(request, org)
        if error:
            return Response({"detail": error}, status=400)

        attendance, assignments, submissions, live, fees, payments, exams, results = scoped_data(
            org, filters
        )

        workbook = Workbook()
        workbook.remove(workbook.active)

        def add_sheet(title, headers, rows):
            ws = workbook.create_sheet(title=title[:31])
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for row in rows:
                ws.append(list(row))
            for column_cells in ws.columns:
                length = max(len(str(cell.value or "")) for cell in column_cells)
                ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 35)
            return ws

        overview_rows = [
            ("Organization", org.name),
            ("Academic Session", filters["session"].name if filters["session"] else "All"),
            ("Class", filters["classroom"].name if filters["classroom"] else "All"),
            ("Section", filters["section"].name if filters["section"] else "All"),
            ("Subject", filters["subject"].name if filters["subject"] else "All"),
            ("Date From", str(filters["date_from"] or "All")),
            ("Date To", str(filters["date_to"] or "All")),
        ]
        att = attendance_summary(attendance)
        expected = decimal_or_zero(fees.aggregate(total=Sum("payable_amount"))["total"])
        collected = decimal_or_zero(payments.aggregate(total=Sum("amount"))["total"])
        overview_rows.extend([
            ("Overall Attendance %", att["overall_percentage"]),
            ("Present", att["present"]),
            ("Absent", att["absent"]),
            ("Late", att["late"]),
            ("Excused", att["excused"]),
            ("Assignments", assignments.count()),
            ("Submissions", submissions.count()),
            ("Live Classes", live.count()),
            ("Published Exams", exams.count()),
            ("Published Results", results.count()),
            ("Expected Fees", float(expected)),
            ("Collected Fees", float(collected)),
            ("Pending Fees", float(max(expected - collected, ZERO))),
        ])
        add_sheet("Overview", ["Metric", "Value"], overview_rows)

        class_rows = []
        class_qs = ClassRoom.objects.filter(organization=org).select_related("academic_session")
        if filters["session"]:
            class_qs = class_qs.filter(academic_session=filters["session"])
        if filters["classroom"]:
            class_qs = class_qs.filter(pk=filters["classroom"].pk)

        for classroom in class_qs:
            enrollments = StudentEnrollment.objects.filter(
                section__classroom=classroom,
                is_active=True,
            )
            class_attendance = attendance.filter(
                attendance_session__section__classroom=classroom
            )
            class_fees = fees.filter(student_id__in=enrollments.values("student_id"))
            class_payments = payments.filter(
                student_fee__student_id__in=enrollments.values("student_id")
            )
            class_expected = decimal_or_zero(
                class_fees.aggregate(total=Sum("payable_amount"))["total"]
            )
            class_collected = decimal_or_zero(
                class_payments.aggregate(total=Sum("amount"))["total"]
            )
            class_rows.append([
                classroom.name,
                classroom.academic_session.name,
                enrollments.count(),
                Section.objects.filter(classroom=classroom).count(),
                attendance_summary(class_attendance)["overall_percentage"],
                assignments.filter(teacher_assignment__section__classroom=classroom).count(),
                results.filter(exam__section__classroom=classroom).count(),
                float(class_expected),
                float(class_collected),
                float(max(class_expected - class_collected, ZERO)),
            ])
        add_sheet(
            "Class-wise Performance",
            ["Class", "Session", "Students", "Sections", "Attendance %", "Assignments",
             "Published Results", "Expected Fees", "Collected Fees", "Pending Fees"],
            class_rows,
        )

        section_rows = []
        section_qs = Section.objects.filter(organization=org).select_related("classroom")
        if filters["classroom"]:
            section_qs = section_qs.filter(classroom=filters["classroom"])
        if filters["section"]:
            section_qs = section_qs.filter(pk=filters["section"].pk)
        for section in section_qs:
            section_rows.append([
                section.classroom.name,
                section.name,
                StudentEnrollment.objects.filter(section=section, is_active=True).count(),
                attendance_summary(attendance.filter(attendance_session__section=section))["overall_percentage"],
            ])
        add_sheet(
            "Section Attendance",
            ["Class", "Section", "Students", "Attendance %"],
            section_rows,
        )

        subject_rows = []
        subject_qs = Subject.objects.filter(organization=org)
        if filters["classroom"]:
            subject_qs = subject_qs.filter(classroom=filters["classroom"])
        if filters["subject"]:
            subject_qs = subject_qs.filter(pk=filters["subject"].pk)
        for subject in subject_qs:
            sr = results.filter(subject=subject)
            total_obtained = decimal_or_zero(sr.aggregate(total=Sum("marks_obtained"))["total"])
            total_max = decimal_or_zero(sr.aggregate(total=Sum("maximum_marks"))["total"])
            average = round(float(total_obtained / total_max * 100), 2) if total_max else 0
            subject_rows.append([
                subject.name,
                attendance_summary(attendance.filter(attendance_session__subject=subject))["overall_percentage"],
                assignments.filter(teacher_assignment__subject=subject).count(),
                sr.count(),
                average,
            ])
        add_sheet(
            "Subject Analytics",
            ["Subject", "Attendance %", "Assignments", "Published Results", "Average %"],
            subject_rows,
        )

        low_attendance_rows = []
        student_ids = attendance.values_list("student_id", flat=True).distinct()
        for enrollment in StudentEnrollment.objects.filter(
            student_id__in=student_ids,
            is_active=True,
        ).select_related("student__user", "section__classroom"):
            summary = attendance_summary(attendance.filter(student=enrollment.student))
            if summary["overall_percentage"] < 75:
                user = enrollment.student.user
                low_attendance_rows.append([
                    user.get_full_name() or user.username,
                    enrollment.roll_number,
                    enrollment.section.classroom.name,
                    enrollment.section.name,
                    summary["present"],
                    summary["absent"],
                    summary["late"],
                    summary["overall_percentage"],
                ])
        add_sheet(
            "Low Attendance",
            ["Student", "Roll Number", "Class", "Section", "Present", "Absent", "Late", "Attendance %"],
            low_attendance_rows,
        )

        teacher_rows = []
        teacher_ids = live.values_list("teacher_assignment__teacher_id", flat=True).distinct()
        for teacher in TeacherProfile.objects.filter(pk__in=teacher_ids).select_related("user"):
            teacher_live = live.filter(teacher_assignment__teacher=teacher)
            teacher_rows.append([
                teacher.user.get_full_name() or teacher.user.username,
                teacher_live.filter(status=LiveClass.Status.SCHEDULED).count(),
                teacher_live.filter(status=LiveClass.Status.COMPLETED).count(),
                teacher_live.filter(status=LiveClass.Status.CANCELLED).count(),
            ])
        add_sheet(
            "Teacher Activity",
            ["Teacher", "Scheduled", "Completed", "Cancelled"],
            teacher_rows,
        )

        outstanding_rows = []
        for fee in fees.select_related("student__user").prefetch_related("payments"):
            paid = fee.paid_amount
            pending = max(fee.payable_amount - paid, ZERO)
            if pending <= ZERO:
                continue
            enrollment = StudentEnrollment.objects.filter(
                student=fee.student,
                is_active=True,
            ).select_related("section__classroom").first()
            outstanding_rows.append([
                fee.student.user.get_full_name() or fee.student.user.username,
                enrollment.roll_number if enrollment else "",
                enrollment.section.classroom.name if enrollment else "",
                enrollment.section.name if enrollment else "",
                float(fee.payable_amount),
                float(paid),
                float(pending),
                fee.calculated_status(),
            ])
        add_sheet(
            "Outstanding Fees",
            ["Student", "Roll Number", "Class", "Section", "Expected", "Paid", "Pending", "Status"],
            outstanding_rows,
        )

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="college-admin-report.xlsx"'
        return response
